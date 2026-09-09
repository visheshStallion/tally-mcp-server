#!/usr/bin/env python3
"""
Converts a Tally Day Book export (saved via the HTTP/XML gateway, even if
saved with a .csv extension - Tally's raw XML response, not real CSV) into
a formatted Excel workbook with three sheets: Day Book, Ledger Entries,
and By Voucher Type.

Usage:
    pip install openpyxl
    python tally_daybook_to_excel.py <input.csv> <output.xlsx> [--company "Company Name"]
"""
import argparse
import csv
import datetime
import re
import sys

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ARIAL = "Arial"
HEADER_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
BODY_FONT = Font(name=ARIAL, size=10)
BOLD_FONT = Font(name=ARIAL, size=10, bold=True)
TITLE_FONT = Font(name=ARIAL, size=14, bold=True, color="1F4E78")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "yyyy-mm-dd"
NUM_FMT = "#,##0.00"


def unescape(s):
    return (
        s.replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&apos;", "'")
    )


def reconstruct_lines(path):
    """Each row in the padded 'csv' is the original raw XML line split on
    every literal comma (no quoting was used when Tally's response was
    saved), padded with trailing empty fields. Re-joining fields with ','
    and stripping only the trailing empty padding recovers the exact
    original line, embedded commas included."""
    lines = []
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        for fields in csv.reader(f):
            while fields and fields[-1] == "":
                fields.pop()
            lines.append(",".join(fields))
    return lines


TAG_RE = re.compile(r'^<([\w.:\-]+)(?:\s[^>]*)?>(.*)</\1>$')
VCH_OPEN_RE = re.compile(r'^<VOUCHER\b')
VCHTYPE_RE = re.compile(r'VCHTYPE="([^"]*)"')


def parse_vouchers(lines):
    vouchers = []
    cur = None
    cur_entry = None
    in_ledger_section = False

    for raw in lines:
        line = raw.strip()
        if VCH_OPEN_RE.match(line):
            m = VCHTYPE_RE.search(line)
            cur = {
                "vchtype": unescape(m.group(1)) if m else "",
                "date": "", "narration": "", "vouchertypename": "",
                "party": "", "vouchernumber": "", "isdeleted": "",
                "amount": "", "entries": [],
            }
            in_ledger_section = False
            continue
        if line == "</VOUCHER>":
            if cur is not None:
                vouchers.append(cur)
            cur = None
            cur_entry = None
            in_ledger_section = False
            continue
        if cur is None:
            continue

        if line in ("<ALLLEDGERENTRIES.LIST>", "<LEDGERENTRIES.LIST>"):
            in_ledger_section = True
            cur_entry = {"ledger": "", "isdeemedpositive": "", "amount": ""}
            continue
        if line in ("</ALLLEDGERENTRIES.LIST>", "</LEDGERENTRIES.LIST>"):
            if cur_entry is not None:
                cur["entries"].append(cur_entry)
            cur_entry = None
            continue

        m = TAG_RE.match(line)
        if not m:
            continue
        tag, content = m.group(1), unescape(m.group(2))

        if cur_entry is not None:
            if tag == "LEDGERNAME" and not cur_entry["ledger"]:
                cur_entry["ledger"] = content
            elif tag == "ISDEEMEDPOSITIVE" and not cur_entry["isdeemedpositive"]:
                cur_entry["isdeemedpositive"] = content
            elif tag == "AMOUNT" and not cur_entry["amount"]:
                cur_entry["amount"] = content
            continue

        if tag == "DATE" and not cur["date"]:
            cur["date"] = content
        elif tag == "NARRATION" and not cur["narration"]:
            cur["narration"] = content
        elif tag == "VOUCHERTYPENAME" and not cur["vouchertypename"]:
            cur["vouchertypename"] = content
        elif tag == "PARTYLEDGERNAME" and not cur["party"]:
            cur["party"] = content
        elif tag == "VOUCHERNUMBER" and not cur["vouchernumber"]:
            cur["vouchernumber"] = content
        elif tag == "ISDELETED" and not cur["isdeleted"]:
            cur["isdeleted"] = content
        elif tag == "AMOUNT" and not in_ledger_section and not cur["amount"]:
            cur["amount"] = content

    return vouchers


def parse_date(s):
    if not s or len(s) != 8:
        return None
    return datetime.date(int(s[0:4]), int(s[4:6]), int(s[6:8]))


def is_plain_number(s):
    if not s:
        return True
    try:
        float(s.strip().replace(",", ""))
        return True
    except ValueError:
        return False


def to_float(s):
    if not s:
        return 0.0
    s = s.strip()
    try:
        return float(s.replace(",", ""))
    except ValueError:
        pass
    # Multi-currency vouchers store e.g. "946836.00 $ @ NGN 1525/ $ = NGN 1443924900.00" --
    # take the converted base-currency (last) figure, respecting its sign.
    m = re.search(r"=\s*(-)?\s*[A-Za-z]{2,5}\s*([\d,]+\.\d+)\s*$", s)
    if m:
        sign = -1 if m.group(1) else 1
        return sign * float(m.group(2).replace(",", ""))
    nums = re.findall(r"-?[\d,]+\.\d+", s)
    if nums:
        return float(nums[-1].replace(",", ""))
    return 0.0


def build_workbook(vouchers, company_name=None):
    for v in vouchers:
        v["_date_obj"] = parse_date(v["date"])
        v["_amount_f"] = to_float(v["amount"])
    vouchers.sort(key=lambda v: (v["_date_obj"] or datetime.date.min, v["vouchernumber"]))

    wb = openpyxl.Workbook()

    # ---------------- Sheet 1: Day Book ----------------
    ws1 = wb.active
    ws1.title = "Day Book"

    header_row = 1
    if company_name:
        ws1.cell(row=1, column=1, value=f"Company: {company_name}").font = TITLE_FONT
        header_row = 3

    headers1 = ["Date", "Voucher Type", "Voucher Number", "Party Ledger", "Narration", "Amount", "Deleted Flag", "Forex Detail (if any)"]
    for col, h in enumerate(headers1, start=1):
        c = ws1.cell(row=header_row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER

    r = header_row + 1
    for v in vouchers:
        vals = [
            v["_date_obj"], v["vouchertypename"] or v["vchtype"], v["vouchernumber"],
            v["party"], v["narration"], v["_amount_f"], v["isdeleted"],
            "" if is_plain_number(v["amount"]) else v["amount"],
        ]
        for col, val in enumerate(vals, start=1):
            c = ws1.cell(row=r, column=col, value=val)
            c.font = BODY_FONT
            c.border = BORDER
            if col == 1:
                c.number_format = DATE_FMT
            if col == 6:
                c.number_format = NUM_FMT
        r += 1

    last_row1 = r - 1
    total_row1 = r
    ws1.cell(row=total_row1, column=5, value="Total").font = BOLD_FONT
    tot_cell = ws1.cell(row=total_row1, column=6, value=f"=SUM(F{header_row+1}:F{last_row1})")
    tot_cell.font = BOLD_FONT
    tot_cell.number_format = NUM_FMT
    ws1.cell(row=total_row1, column=7, value=f'=COUNTA(A{header_row+1}:A{last_row1})&" vouchers"').font = BOLD_FONT

    widths1 = [13, 20, 20, 34, 55, 18, 12, 40]
    for i, w in enumerate(widths1, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = f"A{header_row+1}"

    tab1 = Table(displayName="DayBookTable", ref=f"A{header_row}:H{last_row1}")
    tab1.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws1.add_table(tab1)

    # ---------------- Sheet 2: Ledger Entries ----------------
    ws2 = wb.create_sheet("Ledger Entries")
    headers2 = ["Date", "Voucher Number", "Voucher Type", "Ledger Name", "Dr/Cr", "Amount"]
    for col, h in enumerate(headers2, start=1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

    r = 2
    for v in vouchers:
        for e in v["entries"]:
            drcr = "Dr" if e["isdeemedpositive"] == "Yes" else "Cr"
            amt = abs(to_float(e["amount"]))
            vals = [v["_date_obj"], v["vouchernumber"], v["vouchertypename"] or v["vchtype"], e["ledger"], drcr, amt]
            for col, val in enumerate(vals, start=1):
                c = ws2.cell(row=r, column=col, value=val)
                c.font = BODY_FONT
                c.border = BORDER
                if col == 1:
                    c.number_format = DATE_FMT
                if col == 6:
                    c.number_format = NUM_FMT
            r += 1

    last_row2 = r - 1
    chk_row = r + 1
    ws2.cell(row=chk_row, column=4, value="Total Debit:").font = BOLD_FONT
    d_cell = ws2.cell(row=chk_row, column=6, value=f'=SUMIF(E2:E{last_row2},"Dr",F2:F{last_row2})')
    d_cell.font = BOLD_FONT
    d_cell.number_format = NUM_FMT
    ws2.cell(row=chk_row + 1, column=4, value="Total Credit:").font = BOLD_FONT
    c_cell = ws2.cell(row=chk_row + 1, column=6, value=f'=SUMIF(E2:E{last_row2},"Cr",F2:F{last_row2})')
    c_cell.font = BOLD_FONT
    c_cell.number_format = NUM_FMT
    ws2.cell(row=chk_row + 2, column=4, value="Difference:").font = BOLD_FONT
    diff_cell = ws2.cell(row=chk_row + 2, column=6, value=f"=F{chk_row}-F{chk_row+1}")
    diff_cell.font = BOLD_FONT
    diff_cell.number_format = NUM_FMT

    widths2 = [13, 20, 20, 40, 8, 18]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    tab2 = Table(displayName="LedgerEntriesTable", ref=f"A1:F{last_row2}")
    tab2.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws2.add_table(tab2)

    # ---------------- Sheet 3: By Voucher Type ----------------
    ws3 = wb.create_sheet("By Voucher Type")
    headers3 = ["Voucher Type", "Count", "Total Amount"]
    for col, h in enumerate(headers3, start=1):
        c = ws3.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

    types = sorted({(v["vouchertypename"] or v["vchtype"]) for v in vouchers})
    r = 2
    for t in types:
        ws3.cell(row=r, column=1, value=t).font = BODY_FONT
        cnt_c = ws3.cell(row=r, column=2, value=f"=COUNTIF('Day Book'!B{header_row+1}:B{last_row1},A{r})")
        cnt_c.font = BODY_FONT
        amt_c = ws3.cell(row=r, column=3, value=f"=SUMIF('Day Book'!B{header_row+1}:B{last_row1},A{r},'Day Book'!F{header_row+1}:F{last_row1})")
        amt_c.font = BODY_FONT
        amt_c.number_format = NUM_FMT
        for col in range(1, 4):
            ws3.cell(row=r, column=col).border = BORDER
        r += 1

    last_row3 = r - 1
    ws3.cell(row=r, column=1, value="Total").font = BOLD_FONT
    ws3.cell(row=r, column=2, value=f"=SUM(B2:B{last_row3})").font = BOLD_FONT
    ta = ws3.cell(row=r, column=3, value=f"=SUM(C2:C{last_row3})")
    ta.font = BOLD_FONT
    ta.number_format = NUM_FMT

    widths3 = [24, 10, 20]
    for i, w in enumerate(widths3, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"

    return wb


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input_csv", help="Path to the Tally Day Book export (XML saved as .csv)")
    ap.add_argument("output_xlsx", help="Path to write the formatted Excel workbook")
    ap.add_argument("--company", default=None, help="Company name to show as a title on the Day Book sheet")
    args = ap.parse_args()

    lines = reconstruct_lines(args.input_csv)
    vouchers = parse_vouchers(lines)
    if not vouchers:
        print("No vouchers found - is this really a Tally Day Book export?", file=sys.stderr)
        sys.exit(1)

    wb = build_workbook(vouchers, company_name=args.company)
    wb.save(args.output_xlsx)
    total_entries = sum(len(v["entries"]) for v in vouchers)
    print(f"Wrote {args.output_xlsx}: {len(vouchers)} vouchers, {total_entries} ledger entries.")


if __name__ == "__main__":
    main()

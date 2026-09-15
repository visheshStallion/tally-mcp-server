#!/usr/bin/env python3
"""
Converts the JSON output of the tally-mcp-server "list_vouchers" tool
(as captured via `@modelcontextprotocol/inspector --cli ... --format json`)
into a formatted Excel workbook.

Usage:
    pip install openpyxl
    python tally_vouchers_to_excel.py <input.json> <output.xlsx> [--company "Company Name"]

Where <input.json> is the raw stdout saved from running, e.g.:
    npx @modelcontextprotocol/inspector --cli node dist/index.js \
      --method tools/call --tool-name list_vouchers \
      --tool-arg fromDate=2026-01-01 --tool-arg toDate=2026-12-31 \
      --format json > input.json
"""
import argparse
import datetime
import json
import sys

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ARIAL = "Arial"
HEADER_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
BODY_FONT = Font(name=ARIAL, size=10)
BOLD_FONT = Font(name=ARIAL, size=10, bold=True)
TITLE_FONT = Font(name=ARIAL, size=14, bold=True, color="1F4E78")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "yyyy-mm-dd"
NUM_FMT = "#,##0.00"


def parse_date(s):
    """Tally dates from this tool come back as YYYYMMDD."""
    if not s:
        return None
    digits = str(s).replace("-", "")
    if len(digits) != 8 or not digits.isdigit():
        return s  # leave as-is if it's not the expected format
    return datetime.date(int(digits[0:4]), int(digits[4:6]), int(digits[6:8]))


def unwrap(v):
    """Some tally-mcp-server responses pass through the raw parsed-XML shape
    for fields that carry a TYPE attribute, e.g. {"#text": 123, "@_TYPE":
    "Amount"} instead of just 123, or just {"@_TYPE": "String"} with no
    "#text" key at all when the original XML tag was empty. Unwrap both
    down to a plain value (empty string for the latter case)."""
    if isinstance(v, dict):
        if "#text" in v:
            return v["#text"]
        if "@_TYPE" in v:
            return ""
    return v


def to_float(v):
    v = unwrap(v)
    if v is None:
        return 0.0
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def find_content(obj):
    """Recursively locate a {"content": [...]} envelope, since the tool
    result may or may not be wrapped in an extra {"result": ...} layer
    depending on how it was captured."""
    if isinstance(obj, dict):
        if "content" in obj and isinstance(obj["content"], list):
            return obj
        for v in obj.values():
            found = find_content(v)
            if found is not None:
                return found
    return None


def load_vouchers(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read().strip()

    envelope = None
    for line in raw.splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            candidate = json.loads(line)
        except json.JSONDecodeError:
            continue
        envelope = find_content(candidate)
        if envelope is not None:
            break
    if envelope is None:
        try:
            envelope = find_content(json.loads(raw))
        except json.JSONDecodeError:
            envelope = None
    if envelope is None:
        print("Could not find a JSON tool-result envelope in the input file.", file=sys.stderr)
        sys.exit(1)

    if envelope.get("isError"):
        text = envelope.get("content", [{}])[0].get("text", "")
        print(f"The tool call itself returned an error: {text}", file=sys.stderr)
        sys.exit(1)

    inner_text = envelope["content"][0]["text"]
    payload = json.loads(inner_text)
    vouchers = payload.get("vouchers", [])
    for v in vouchers:
        for key in ("date", "voucherType", "voucherNumber", "partyLedger", "narration", "amount"):
            if key in v:
                v[key] = unwrap(v[key])
    return vouchers


def build_workbook(vouchers, company_name=None):
    for v in vouchers:
        v["_date_obj"] = parse_date(v.get("date"))
        v["_amount_f"] = to_float(v.get("amount"))
    vouchers.sort(key=lambda v: (v["_date_obj"] if isinstance(v["_date_obj"], datetime.date) else datetime.date.min, str(v.get("voucherNumber") or "")))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vouchers"

    header_row = 1
    if company_name:
        ws.cell(row=1, column=1, value=f"Company: {company_name}").font = TITLE_FONT
        header_row = 3

    headers = ["Date", "Voucher Type", "Voucher Number", "Party Ledger", "Narration", "Amount"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER

    r = header_row + 1
    for v in vouchers:
        vals = [
            v["_date_obj"], v.get("voucherType"), v.get("voucherNumber"),
            v.get("partyLedger"), v.get("narration"), v["_amount_f"],
        ]
        for col, val in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = BODY_FONT
            c.border = BORDER
            if col == 1 and isinstance(val, datetime.date):
                c.number_format = DATE_FMT
            if col == 6:
                c.number_format = NUM_FMT
        r += 1

    last_row = r - 1
    total_row = r
    ws.cell(row=total_row, column=5, value="Total").font = BOLD_FONT
    tot_cell = ws.cell(row=total_row, column=6, value=f"=SUM(F{header_row+1}:F{last_row})")
    tot_cell.font = BOLD_FONT
    tot_cell.number_format = NUM_FMT
    ws.cell(row=total_row, column=7, value=f'=COUNTA(A{header_row+1}:A{last_row})&" vouchers"').font = BOLD_FONT

    widths = [13, 20, 20, 34, 55, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row+1}"

    if last_row >= header_row:
        tab = Table(displayName="VouchersTable", ref=f"A{header_row}:F{last_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)

    return wb


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input_json", help="Path to the saved list_vouchers tool output")
    ap.add_argument("output_xlsx", help="Path to write the formatted Excel workbook")
    ap.add_argument("--company", default=None, help="Company name to show as a title")
    args = ap.parse_args()

    vouchers = load_vouchers(args.input_json)
    if not vouchers:
        print("No vouchers found in the input file.", file=sys.stderr)
        sys.exit(1)

    wb = build_workbook(vouchers, company_name=args.company)
    wb.save(args.output_xlsx)
    print(f"Wrote {args.output_xlsx}: {len(vouchers)} vouchers.")


if __name__ == "__main__":
    main()
